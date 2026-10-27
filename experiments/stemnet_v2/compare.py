#!/usr/bin/env python3
"""StemNet v2: the base-only head and the base + tree head side by side,
against the survey bases, at every threshold.

For each checkpoint in V2_COMPARE (default head_best.pth = base-only,
head_tree_best.pth = base + tree):
  1. val and test predicted once (saved and reused, as v2_validate does; a
     checkpoint whose predictions are already saved is not predicted again)
  2. for every gate threshold in V2_CURVE_P (0.3 ... 0.9): detections with
     vote mass >= V2_CURVE_MIN_MASS (1), peaks V2_CURVE_SEP (0.3 m) apart,
     ranked by vote mass, matched one-to-one to the survey bases (greedy by
     vote mass) at 0.5 m and 0.3 m: TP = matched, FP = any other detection
     (an unmatched one outside the split's cells is ignored, the survey stops
     there). -> precision-recall curve, AP, FROC (recall at 1-50 false per ha)
  3. chosen on val, per model, nothing on test: the gate threshold with the best
     AP@0.5, then the vote-mass cut with the best F1 on that val curve. Test is
     reported at that point, plus over its whole curve.
TP / FP are against the survey only: unsurveyed trees count as false, so every
precision here is a lower bound (see v2_inspect's close-ups).

Outputs:
  results/stemnet_v2_compare.xlsx   sheets:
      README       what every column means
      Summary      one row per model: AP, CPM, max recall, the val-chosen point
                   on test, recall at fixed precision, precision at fixed recall
      PerGate      every model x split x gate threshold: AP, CPM, max recall
      Breakdown    recall by survey tier and trunk diameter, per model, at the
                   val-chosen point and at maximum recall
      Curves       every rank of the val and test curves at each model's chosen
                   gate: vote mass, TP, FP, precision, recall, FP / ha
      Detections   every detection at the chosen gate: position (frame, UTM),
                   vote mass, rank, matched survey base, distance to the
                   nearest survey base, kept at the val-chosen point
      SurveyTrees  every val / test survey base: tier, DBH, and per model found
                   at the val-chosen point / at maximum recall, distance to and
                   vote mass of the nearest detection
  results/stemnet_v2_compare.txt    the summary in text
  results/58_stemnet_v2_compare.png PR and FROC curves of both models, AP by gate
"""
from __future__ import annotations

import csv
import json
import os
import sys
import time
from pathlib import Path

import numpy as np
import torch

sys.path.insert(0, str(Path(__file__).resolve().parent))
from core import OUT, RESULTS  # noqa: E402
import curves as C  # noqa: E402
import data as D  # noqa: E402
import infer as I  # noqa: E402
import train as T  # noqa: E402
import validate as Va  # noqa: E402
import vote as V  # noqa: E402

CKPTS = [c.strip() for c in os.environ.get("V2_COMPARE", "head_best.pth,head_tree_best.pth").split(",") if c.strip()]
P_GRID = C.P_GRID
MIN_MASS = C.MIN_MASS
SEP = float(os.environ.get("V2_CURVE_SEP", "0.3"))
PREC_AT = (0.9, 0.8, 0.7)
REC_AT = (0.7, 0.8, 0.9)
DBH_BINS = Va.DBH_BINS
SPLITS = {1: "val", 2: "test"}
LOG = RESULTS / "stemnet_v2_compare.txt"


def log(msg):
    line = f"[v2-compare] {msg}"
    print(line, flush=True)
    with open(LOG, "a") as f:
        f.write(line + "\n")


def label_of(ck):
    return "base+tree" if ck.get("task") == "base+tree" else "base-only"


def run_all(regions, G):
    """Per checkpoint: saved val / test predictions (predicting what is missing), then
    its curves; one model's predictions in memory at a time. -> (list of model dicts, ha)."""
    models, enc = [], None
    store = D.TileStore(keep=I.KEEP)
    Va.log = log                                               # validate's helpers log here
    for name in CKPTS:
        path = T.MODELS / name
        if not path.exists():
            log(f"!! {name} not found in {T.MODELS} — left out")
            continue
        if enc is None:
            enc = M_encoder()
        runner, ck, cpath = I.load_runner(log, name, enc)
        R = {}
        for sid in SPLITS:
            R[sid], _ = Va.predict_split(runner, store, regions, sid, np.random.default_rng(T.SEED), cpath, ck)
            R[sid] = {k: R[sid][k] for k in ("xy", "p", "off")}
        del runner
        if torch.cuda.is_available():
            torch.cuda.empty_cache()
        md = dict(name=cpath.stem, label=label_of(ck), step=int(ck["step"]), R=R)
        ha = evaluate(md, regions, G)
        models.append(md)
    if not models:
        raise SystemExit("[v2-compare] no checkpoint to compare — train first (v2_train, v2_train_tree)")
    return models, ha


def M_encoder():
    import model as M
    return M.load_sonata(T.DEVICE) if T.ENCODER == "sonata" else M.FakeEncoder()


def survey(regions):
    rows = list(csv.DictReader(open(OUT / "label_bases.csv")))
    xy = np.array([[float(r["x"]), float(r["y"])] for r in rows]).reshape(-1, 2)
    sp = C.split_at(regions["split"], xy)
    out = {}
    for sid in SPLITS:
        m = sp == sid
        col = lambda k, f=str: np.array([f(r[k]) for r in rows])[m]  # noqa: E731
        out[sid] = dict(xy=xy[m], id=col("base", int), tier=col("tier"), r=col("radius", float),
                        measured=col("radius_measured") == "1", has_points=col("has_points") == "1")
    return out


def point_stats(c, n_gt, ha):
    """From a curve: best F1 (and its vote-mass cut), recall at fixed precision, precision at fixed recall."""
    if not len(c["tp"]):
        return dict(best_f1=0.0, best_f1_mass=np.inf, rec_at={p: 0.0 for p in PREC_AT}, prec_at={r: 0.0 for r in REC_AT})
    rank = np.arange(1, len(c["tp"]) + 1)
    prec, rec = c["tp"] / rank, c["tp"] / max(n_gt, 1)
    f1 = 2 * c["tp"] / (rank + n_gt)
    k = int(np.argmax(f1))
    return dict(best_f1=float(f1[k]), best_f1_mass=float(c["score"][k]),
                rec_at={p: float(rec[prec >= p].max()) if (prec >= p).any() else 0.0 for p in PREC_AT},
                prec_at={r: float(prec[rec >= r].max()) if (rec >= r).any() else 0.0 for r in REC_AT})


def found_flags(det, sc, G, regions, sid, mass):
    k = sc >= mass
    ign = C.split_at(regions["split"], det[k]) != sid if k.any() else None
    m = V.match(det[k], sc[k], G["xy"], 0.5, ignore=ign)
    f = np.zeros(len(G["xy"]), bool)
    f[m["pair"][m["pair"] >= 0]] = True
    return f


def breakdown(G, found):
    out = {}
    for t in ("A", "B", "C"):
        m = G["tier"] == t
        if m.any():
            out[f"tier {t}"] = (float(found[m].mean()), int(m.sum()))
    dbh = 2 * G["r"]
    for lo, hi in zip(DBH_BINS[:-1], DBH_BINS[1:]):
        m = G["measured"] & (dbh >= lo) & (dbh < hi)
        if m.any():
            out[f"DBH {lo * 100:.0f}-{hi * 100:.0f} cm" if np.isfinite(hi) else f"DBH >= {lo * 100:.0f} cm"] = \
                (float(found[m].mean()), int(m.sum()))
    out["all"] = (float(found.mean()) if len(found) else 0.0, int(len(found)))
    return out


def evaluate(md, regions, G):
    occ, split = regions["occupancy"], regions["split"]
    ha = {sid: float(((split == sid) & (occ > 0)).sum()) / 1e4 for sid in SPLITS}
    t0 = time.time()
    md["gates"] = {}
    for sid in SPLITS:
        xy = np.asarray(md["R"][sid]["xy"], np.float64)
        p, off = np.asarray(md["R"][sid]["p"]), np.asarray(md["R"][sid]["off"], np.float64)
        for pt in P_GRID:
            det, sc, nv = V.detect(xy, p, off, p_thresh=pt, peak_sep=SEP, peak_min=MIN_MASS)
            ign = C.split_at(split, det) != sid if len(det) else None
            c5, c3 = C.curve(det, sc, G[sid]["xy"], 0.5, ign), C.curve(det, sc, G[sid]["xy"], 0.3, ign)
            md["gates"][(sid, pt)] = dict(det=det, sc=sc, nv=nv, c5=c5, c3=c3, ap5=C.ap(c5), ap3=C.ap(c3),
                                          froc=C.froc(c5, ha[sid]),
                                          max_rec=float(c5["tp"][-1] / max(c5["n_gt"], 1)) if len(c5["tp"]) else 0.0,
                                          stats=point_stats(c5, c5["n_gt"], ha[sid]))
    md.pop("R")                                             # predictions no longer needed
    md["gate"] = max(P_GRID, key=lambda pt: md["gates"][(1, pt)]["ap5"])
    md["mass"] = md["gates"][(1, md["gate"])]["stats"]["best_f1_mass"]
    g = md["gates"][(2, md["gate"])]
    md["test_point"] = C.at_mass(g["c5"], md["mass"])
    md["found"] = {sid: dict(chosen=found_flags(md["gates"][(sid, md["gate"])]["det"], md["gates"][(sid, md["gate"])]["sc"],
                                                G[sid], regions, sid, md["mass"]),
                             maxrec=found_flags(md["gates"][(sid, md["gate"])]["det"], md["gates"][(sid, md["gate"])]["sc"],
                                                G[sid], regions, sid, MIN_MASS))
                   for sid in SPLITS}
    log(f"{md['label']} ({md['name']}, step {md['step']}): curves for {len(P_GRID)} gate thresholds x 2 splits "
        f"in {(time.time() - t0) / 60:.1f} min")
    return ha


def report(models, G, ha):
    log(f"survey bases: val {len(G[1]['xy'])}, test {len(G[2]['xy'])} | scanned ha: val {ha[1]:.2f}, test {ha[2]:.2f} | "
        "TP / FP against the survey only (unsurveyed trees count as false: precision is a lower bound)")
    for md in models:
        log(f"{md['label']} ({md['name']}, step {md['step']}): AP@0.5 by gate threshold, val / test:")
        for pt in P_GRID:
            v, t = md["gates"][(1, pt)], md["gates"][(2, pt)]
            log(f"  gate >= {pt:.1f}: val AP {v['ap5']:.3f} | test AP@0.5 {t['ap5']:.3f} AP@0.3 {t['ap3']:.3f} "
                f"CPM {np.mean(t['froc']):.3f} max recall {t['max_rec']:.3f} ({len(t['det'])} det)")
    log("=" * 100)
    log("TEST (held out), each model at its own val-chosen gate threshold:")
    for md in models:
        g = md["gates"][(2, md["gate"])]
        pr, rc = md["test_point"]
        st = g["stats"]
        log(f"{md['label']:<10} gate >= {md['gate']:.1f}: AP@0.5 {g['ap5']:.3f} | AP@0.3 {g['ap3']:.3f} | "
            f"CPM {np.mean(g['froc']):.3f} | max recall {g['max_rec']:.3f}")
        log(f"{'':<10} val-chosen cut (vote mass >= {md['mass']:.0f}): precision {pr:.3f} recall {rc:.3f} "
            f"F1 {2 * pr * rc / max(pr + rc, 1e-9):.3f} | best F1 anywhere on the test curve {st['best_f1']:.3f} "
            f"(oracle, not a result)")
        log(f"{'':<10} recall at precision >= " + ", ".join(f"{p:g}: {st['rec_at'][p]:.3f}" for p in PREC_AT)
            + " | precision at recall >= " + ", ".join(f"{r:g}: {st['prec_at'][r]:.3f}" for r in REC_AT))
        log(f"{'':<10} FROC recall at " + " / ".join(f"{f}" for f in C.FP_HA) + " false per ha: "
            + " / ".join(f"{r:.2f}" for r in g["froc"]))
        for kind in ("chosen", "maxrec"):
            b = breakdown(G[2], md["found"][2][kind])
            log(f"{'':<10} recall {'at the val-chosen cut' if kind == 'chosen' else 'at maximum recall'}: "
                + " | ".join(f"{k} {v:.2f} of {n}" for k, (v, n) in b.items()))


def write_xlsx(models, G, ha, regions, origin):
    from openpyxl import Workbook
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("README")
    for line in (__doc__ or "").splitlines():
        ws.append([line])
    ws = wb.create_sheet("Summary")
    ws.append(["model", "checkpoint", "step", "gate (val-chosen)", "vote mass cut (val-chosen)",
               "val AP@0.5", "test AP@0.5", "test AP@0.3", "test CPM", "test max recall",
               "test precision @cut", "test recall @cut", "test F1 @cut", "test best F1 on curve (oracle)"]
              + [f"test recall @ precision>={p:g}" for p in PREC_AT] + [f"test precision @ recall>={r:g}" for r in REC_AT]
              + [f"test recall @ {f} FP/ha" for f in C.FP_HA])
    for md in models:
        g, v = md["gates"][(2, md["gate"])], md["gates"][(1, md["gate"])]
        pr, rc = md["test_point"]
        ws.append([md["label"], md["name"], md["step"], md["gate"], round(md["mass"], 2), round(v["ap5"], 4),
                   round(g["ap5"], 4), round(g["ap3"], 4), round(float(np.mean(g["froc"])), 4), round(g["max_rec"], 4),
                   round(pr, 4), round(rc, 4), round(2 * pr * rc / max(pr + rc, 1e-9), 4), round(g["stats"]["best_f1"], 4)]
                  + [round(g["stats"]["rec_at"][p], 4) for p in PREC_AT] + [round(g["stats"]["prec_at"][r], 4) for r in REC_AT]
                  + [round(r, 4) for r in g["froc"]])
    ws = wb.create_sheet("PerGate")
    ws.append(["model", "split", "gate", "detections", "AP@0.5", "AP@0.3", "CPM", "max recall", "best F1 on curve"]
              + [f"recall @ {f} FP/ha" for f in C.FP_HA])
    for md in models:
        for sid in SPLITS:
            for pt in P_GRID:
                g = md["gates"][(sid, pt)]
                ws.append([md["label"], SPLITS[sid], pt, len(g["det"]), round(g["ap5"], 4), round(g["ap3"], 4),
                           round(float(np.mean(g["froc"])), 4), round(g["max_rec"], 4), round(g["stats"]["best_f1"], 4)]
                          + [round(r, 4) for r in g["froc"]])
    ws = wb.create_sheet("Breakdown")
    ws.append(["model", "split", "point", "group", "recall@0.5", "survey bases"])
    for md in models:
        for sid in SPLITS:
            for kind, nm in (("chosen", "val-chosen cut"), ("maxrec", "maximum recall")):
                for k, (v, n) in breakdown(G[sid], md["found"][sid][kind]).items():
                    ws.append([md["label"], SPLITS[sid], nm, k, round(v, 4), n])
    ws = wb.create_sheet("Curves")
    ws.append(["model", "split", "gate", "rank", "vote mass", "TP", "FP", "precision", "recall", "FP per ha", "TP @0.3 m", "FP @0.3 m"])
    for md in models:
        for sid in SPLITS:
            g = md["gates"][(sid, md["gate"])]
            c5, c3 = g["c5"], g["c3"]
            for k in range(len(c5["tp"])):
                ws.append([md["label"], SPLITS[sid], md["gate"], k + 1, round(float(c5["score"][k]), 2), int(c5["tp"][k]),
                           int(c5["fp"][k]), round(float(c5["tp"][k] / (k + 1)), 4), round(float(c5["tp"][k] / max(c5["n_gt"], 1)), 4),
                           round(float(c5["fp"][k] / ha[sid]), 3),
                           int(c3["tp"][k]) if k < len(c3["tp"]) else None, int(c3["fp"][k]) if k < len(c3["fp"]) else None])
    ws = wb.create_sheet("Detections")
    ws.append(["model", "split", "gate", "rank", "x", "y", "utm_x", "utm_y", "vote mass", "votes",
               "matched survey base (-1 false, -2 ignored)", "distance to nearest survey base (m)", "kept at val-chosen cut"])
    from scipy.spatial import cKDTree
    for md in models:
        for sid in SPLITS:
            g = md["gates"][(sid, md["gate"])]
            det, sc, nv = g["det"], g["sc"], g["nv"]
            if not len(det):
                continue
            ign = C.split_at(regions["split"], det) != sid
            pair = V.match(det, sc, G[sid]["xy"], 0.5, ignore=ign)["pair"]
            dist = cKDTree(G[sid]["xy"]).query(det)[0] if len(G[sid]["xy"]) else np.full(len(det), np.inf)
            o = np.argsort(-sc, kind="stable")
            for k, i in enumerate(o):
                ws.append([md["label"], SPLITS[sid], md["gate"], k + 1, round(float(det[i, 0]), 3), round(float(det[i, 1]), 3),
                           round(float(det[i, 0] + origin[0]), 3), round(float(det[i, 1] + origin[1]), 3), round(float(sc[i]), 2),
                           int(nv[i]), int(G[sid]["id"][pair[i]]) if pair[i] >= 0 else int(pair[i]), round(float(dist[i]), 3),
                           int(sc[i] >= md["mass"])])
    ws = wb.create_sheet("SurveyTrees")
    head = ["split", "survey base", "x", "y", "utm_x", "utm_y", "tier", "radius measured", "DBH (cm)", "scan points under it"]
    for md in models:
        head += [f"{md['label']}: found @cut", f"{md['label']}: found @max recall",
                 f"{md['label']}: nearest detection (m)", f"{md['label']}: its vote mass"]
    ws.append(head)
    for sid in SPLITS:
        Gs = G[sid]
        near = []
        for md in models:
            g = md["gates"][(sid, md["gate"])]
            if len(g["det"]):
                d, i = cKDTree(g["det"]).query(Gs["xy"])
                near.append((d, g["sc"][i]))
            else:
                near.append((np.full(len(Gs["xy"]), np.inf), np.zeros(len(Gs["xy"]))))
        for j in range(len(Gs["xy"])):
            row = [SPLITS[sid], int(Gs["id"][j]), round(float(Gs["xy"][j, 0]), 3), round(float(Gs["xy"][j, 1]), 3),
                   round(float(Gs["xy"][j, 0] + origin[0]), 3), round(float(Gs["xy"][j, 1] + origin[1]), 3), str(Gs["tier"][j]),
                   int(Gs["measured"][j]), round(float(2 * Gs["r"][j] * 100), 1) if Gs["measured"][j] else None,
                   int(Gs["has_points"][j])]
            for md, (d, s) in zip(models, near):
                row += [int(md["found"][sid]["chosen"][j]), int(md["found"][sid]["maxrec"][j]),
                        round(float(d[j]), 3) if np.isfinite(d[j]) else None, round(float(s[j]), 2)]
            ws.append(row)
    wb.save(RESULTS / "stemnet_v2_compare.xlsx")


def figure(models):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    cols = ["tab:blue", "tab:green", "tab:purple", "tab:orange"]
    fig, axs = plt.subplots(1, 3, figsize=(21, 6.5))
    for md, col in zip(models, cols):
        for sid, ls in ((1, "--"), (2, "-")):
            g = md["gates"][(sid, md["gate"])]
            c = g["c5"]
            if len(c["tp"]):
                rank = np.arange(1, len(c["tp"]) + 1)
                axs[0].plot(c["tp"] / max(c["n_gt"], 1), c["tp"] / rank, ls, c=col,
                            label=f"{md['label']} {SPLITS[sid]}: AP {g['ap5']:.3f}")
                axs[1].plot(np.maximum(c["fp"], 0.5) / md.get("ha", {}).get(sid, 1.0), c["tp"] / max(c["n_gt"], 1), ls, c=col,
                            label=f"{md['label']} {SPLITS[sid]}: CPM {np.mean(g['froc']):.3f}")
        pr, rc = md["test_point"]
        axs[0].plot(rc, pr, "*", c=col, ms=14, mec="k")
        axs[2].plot(P_GRID, [md["gates"][(1, pt)]["ap5"] for pt in P_GRID], "o--", c=col, label=f"{md['label']} val")
        axs[2].plot(P_GRID, [md["gates"][(2, pt)]["ap5"] for pt in P_GRID], "o-", c=col, label=f"{md['label']} test")
    axs[0].set_xlabel("recall (share of survey bases found)"); axs[0].set_ylabel("precision vs survey (lower bound)")
    axs[0].set_xlim(0, 1); axs[0].set_ylim(0, 1.02); axs[0].legend(loc="lower left", fontsize=8)
    axs[0].set_title("precision-recall @0.5 m, each model at its val-chosen gate (* = val-chosen cut on test)")
    axs[1].set_xscale("log"); axs[1].set_xlabel("false detections per hectare (vs survey)"); axs[1].set_ylabel("recall @0.5 m")
    axs[1].set_ylim(0, 1.02); axs[1].legend(loc="upper left", fontsize=8); axs[1].set_title("FROC")
    for f in C.FP_HA:
        axs[1].axvline(f, c="grey", lw=0.5, ls=":")
    axs[2].set_xlabel("gate threshold"); axs[2].set_ylabel("AP@0.5"); axs[2].legend(fontsize=8)
    axs[2].set_title("AP by gate threshold (dashed val, solid test)")
    fig.tight_layout(); fig.savefig(RESULTS / "58_stemnet_v2_compare.png", dpi=110); plt.close(fig)


def main():
    RESULTS.mkdir(parents=True, exist_ok=True)
    LOG.write_text("")
    frame = json.loads((OUT / "frame.json").read_text())
    origin = np.asarray(frame["origin_m"], np.float64)
    regions = np.load(OUT / "regions.npz")
    log(f"comparing {', '.join(CKPTS)} against the survey bases | gate thresholds {P_GRID} | "
        f"vote mass >= {MIN_MASS:g}, peaks {SEP:g} m apart")
    G = survey(regions)
    models, ha = run_all(regions, G)
    for md in models:
        md["ha"] = ha
    report(models, G, ha)
    write_xlsx(models, G, ha, regions, origin)
    figure(models)
    log("-> results/stemnet_v2_compare.xlsx, stemnet_v2_compare.txt, 58_stemnet_v2_compare.png")


if __name__ == "__main__":
    main()
