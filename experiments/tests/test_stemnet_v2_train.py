#!/usr/bin/env python3
"""StemNet v2 training, on CPU with a stand-in encoder: the pieces that can be
silently wrong without a GPU ever noticing.

The scene is built through the real prep stages (features, labels) from
synthetic grid tiles, so the trainer reads exactly what it will read on the
campus."""
from __future__ import annotations

import csv
import json
import os
import subprocess
import sys
import tempfile
from pathlib import Path

import numpy as np

HERE = Path(__file__).resolve().parents[1]
V2 = HERE / "stemnet_v2"
TILE, W = 10.0, 30.0


def _build(tmp: Path):
    rng = np.random.default_rng(1)
    pos = []
    while len(pos) < 40:
        p = rng.uniform(2, W - 2, 2)
        if all(np.hypot(*(p - q)) >= 1.2 for q in pos):
            pos.append(p)
    pos = np.array(pos); rad = rng.uniform(0.05, 0.15, len(pos))
    g = np.arange(0.02, W, 0.04)
    gx, gy = np.meshgrid(g, g)
    parts = [np.column_stack([gx.ravel(), gy.ravel(), 0.03 * gx.ravel()])]
    for (x, y), r in zip(pos, rad):
        a = np.linspace(0, 2 * np.pi, max(int(2 * np.pi * r / 0.02), 8), endpoint=False)
        z = np.arange(0, 4.0, 0.02)
        A, Z = np.meshgrid(a, z)
        parts.append(np.column_stack([x + r * np.cos(A.ravel()), y + r * np.sin(A.ravel()), 0.03 * x + Z.ravel()]))
    P = np.concatenate(parts).astype(np.float32)
    d = tmp / "data" / "stemnet_v2"
    (d / "grid0.02").mkdir(parents=True)
    ti = np.floor(P[:, 0] / TILE).astype(int); tj = np.floor(P[:, 1] / TILE).astype(int)
    for i, j in sorted(set(zip(ti, tj))):
        m = (ti == i) & (tj == j)
        np.savez(d / "grid0.02" / f"tile_{i}_{j}.npz", xyz=P[m], rgb=np.zeros((m.sum(), 3), np.uint8),
                 intensity=np.zeros(m.sum(), np.uint8))
    with open(d / "bases.csv", "w", newline="") as f:
        w = csv.writer(f); w.writerow(["id", "x", "y", "radius", "base_tier", "fork_group"])
        for k, ((x, y), r) in enumerate(zip(pos, rad)):
            w.writerow([f"1-{k}", x, y, r, "A", k])
    split = np.full((int(W), int(W)), 0, np.uint8)
    split[:, 18:20] = 3; split[:, 20:] = 1                  # train | buffer | val
    np.savez(d / "regions.npz", split=split, occupancy=np.ones_like(split, np.int32))
    (d / "frame.json").write_text(json.dumps(dict(grids={"0.02": {"stamp": {}}}, bases={}, splits={},
                                                  origin_m=[0, 0, 0], tile_m=TILE)))
    env = dict(os.environ, DATA_DIR=str(tmp / "data"), RESULTS_DIR=str(tmp / "res"), V2_TILE=str(TILE))
    for stage in ("prep_features.py", "prep_labels.py", "prep_cache.py"):
        r = subprocess.run([sys.executable, str(V2 / stage)], env=env, capture_output=True, text=True)
        assert r.returncode == 0, (stage, r.stdout[-2000:], r.stderr[-2000:])
    return env, pos, rad


def _with_env(env):
    for k, v in env.items():
        os.environ[k] = v
    sys.path.insert(0, str(V2))
    for m in ("core", "data", "model", "vote", "train", "infer", "validate"):
        sys.modules.pop(m, None)


def test_crop_equals_brute_force_and_augmentation_keeps_bases():
    with tempfile.TemporaryDirectory() as t:
        env, pos, rad = _build(Path(t))
        _with_env(env)
        import data as D
        store = D.TileStore()
        allxyz = np.concatenate([store.tile(ij)["xyz"] for ij in sorted(store.have)])
        c = store.crop(10.0, 10.0, 8.0)                      # straddles four tiles
        want = allxyz[(np.abs(allxyz[:, 0] - 10) <= 4) & (np.abs(allxyz[:, 1] - 10) <= 4)]
        assert len(c["xyz"]) == len(want)
        assert c["gate"].sum() > 100
        base_before = c["xyz"][:, :2] + c["offset"]
        rng = np.random.default_rng(3)
        D.AUG_P, D.JIT_P = 1.0, 1.0                           # every augmentation, every time
        A = D.augment(c, rng, np.array([10.0, 10.0]))
        k = c["gate"] == 1
        want_base = (base_before[k] - [10, 10]) @ A.T + [10, 10]
        got_base = c["xyz"][k, :2] + c["offset"][k]
        assert np.abs(got_base - want_base).max() < 1e-3, np.abs(got_base - want_base).max()
        assert np.allclose(np.linalg.norm(c["normal"], axis=1), 1, atol=1e-4)


def test_cache_crops_equal_tile_crops():
    with tempfile.TemporaryDirectory() as t:
        env, pos, rad = _build(Path(t))
        _with_env(env)
        import data as D
        a, b = D.TileStore(), D.CacheStore()
        assert type(D.open_store(lambda m: None)) is D.CacheStore
        rng = np.random.default_rng(7)
        for cx, cy in [(10.0, 10.0), (24.9, 5.0), (3.0, 27.5)] + [tuple(v) for v in rng.uniform(0, W, (12, 2))]:
            ca, cb = a.crop(cx, cy, 8.0), b.crop(cx, cy, 8.0)
            if ca is None:
                assert cb is None; continue
            oa = np.lexsort(ca["xyz"].T); ob = np.lexsort(cb["xyz"].T)
            assert len(oa) == len(ob), (cx, cy, len(oa), len(ob))
            for k in ("xyz", "feat", "gate", "offset", "sem", "normal", "tile"):
                assert np.array_equal(ca[k][oa], cb[k][ob]), (cx, cy, k)
            t_ = cb["tile"][0]
            xyz = b.tile((int(t_) // 100000, int(t_) % 100000))["xyz"]
            m = cb["tile"] == t_
            assert np.array_equal(np.asarray(xyz)[cb["row"][m]].astype(np.float64), cb["xyz"][m])


def test_inference_path_finds_every_base_from_perfect_predictions():
    """Feed the true gate and offset through predict_region -> detect -> match:
    whatever the model learns, this path must not lose bases on its own."""
    with tempfile.TemporaryDirectory() as t:
        env, pos, rad = _build(Path(t))
        _with_env(env)
        import data as D, vote as V
        store = D.TileStore()
        oracle = lambda c: (c["gate"].astype(np.float32), c["offset"].astype(np.float32))  # noqa: E731
        bbox = (20.0, 30.0, 0.0, 30.0)
        xy, p, off = V.predict_region(oracle, store, bbox, 300000, np.random.default_rng(0))
        det, score, nv = V.detect(xy, p, off)
        gt = pos[(pos[:, 0] >= 20)]
        m = V.match(det, score, gt, 0.3, bbox)
        assert m["recall"] == 1.0 and m["precision"] == 1.0, {k: v for k, v in m.items() if k != "pair"}
        assert m["rmse"] < 0.03, m["rmse"]
        rng = np.random.default_rng(5)                         # an imperfect model

        def noisy(c):
            g = c["gate"].astype(np.float32)
            g[(g == 0) & (rng.random(len(g)) < 0.05)] = 0.9
            return g, (c["offset"] + rng.normal(0, 0.05, c["offset"].shape)).astype(np.float32)
        xy, p, off = V.predict_region(noisy, store, bbox, 300000, rng)
        det, score, nv = V.detect(xy, p, off)
        m = V.match(det, score, gt, 0.3, bbox)
        assert m["recall"] == 1.0 and m["rmse"] < 0.05, {k: v for k, v in m.items() if k != "pair"}


def test_training_runs_evaluates_saves_and_resumes():
    import torch
    with tempfile.TemporaryDirectory() as t:
        env, pos, rad = _build(Path(t))
        for q in sorted((Path(t) / "data" / "stemnet_v2" / "feat0.02").glob("tile_*.npz")):
            d = dict(np.load(q))                              # run 1 died on this: an isolated
            d["feat"][::50, 6] = np.inf                       # point overflowing float16
            np.savez(q, **d)
        fj = Path(t) / "data" / "stemnet_v2" / "frame.json"
        fr = json.loads(fj.read_text()); fr.pop("cache"); fj.write_text(json.dumps(fr))
        r = subprocess.run([sys.executable, str(V2 / "prep_cache.py")], env=env, capture_output=True, text=True)
        assert r.returncode == 0, r.stdout[-1500:] + r.stderr[-1500:]
        assert json.loads(fj.read_text())["cache"]["feature_values_saturated"] > 0
        env.update(V2_ENCODER="fake", V2_DEVICE="cpu", V2_STEPS="12", V2_EVAL_EVERY="6", V2_ACCUM="1",
                   V2_CAP="40000", V2_PROGRESS_S="0")
        r = subprocess.run([sys.executable, str(V2 / "train.py")], env=env, capture_output=True, text=True)
        assert r.returncode == 0, r.stdout[-3000:] + r.stderr[-3000:]
        log = (Path(t) / "res" / "stemnet_v2_train.txt").read_text()
        assert log.count("EVAL step") == 2 and "done: 12 steps" in log, log[-1500:]
        assert "reading crops from the sorted cache" in log, log[:1500]
        assert (Path(t) / "data" / "stemnet_v2" / "models" / "head_last.pth").exists()
        assert "non-finite" not in log and "not applied" not in log, log[-1500:]
        models = Path(t) / "data" / "stemnet_v2" / "models"
        ck = torch.load(models / "head_last.pth", weights_only=False)
        assert all(bool(torch.isfinite(v).all()) for v in ck["head"].values())
        env["V2_STEPS"] = "18"
        r = subprocess.run([sys.executable, str(V2 / "train.py")], env=env, capture_output=True, text=True)
        assert r.returncode == 0, r.stdout[-3000:] + r.stderr[-3000:]
        assert "resuming from step 12" in r.stdout, r.stdout[-1500:]
        best_step = torch.load(models / "head_best.pth", weights_only=False)["step"]
        ck = torch.load(models / "head_last.pth", weights_only=False)   # a NaN'd latest state
        ck["head"] = {k: v * float("nan") for k, v in ck["head"].items()}
        torch.save(ck, models / "head_last.pth")
        env["V2_STEPS"] = "20"
        r = subprocess.run([sys.executable, str(V2 / "train.py")], env=env, capture_output=True, text=True)
        assert r.returncode == 0, r.stdout[-3000:] + r.stderr[-3000:]
        assert "head_last.pth (step 18) has non-finite weights" in r.stdout, r.stdout[-1500:]
        assert f"resuming from step {best_step}" in r.stdout, r.stdout[-1500:]


def _add_sem(t, pos, rad):
    """Tree / non-tree labels as prep_semantic writes them: trunk = 1, ground = 2."""
    d = Path(t) / "data" / "stemnet_v2"
    (d / "sem0.02").mkdir()
    for q in sorted((d / "grid0.02").glob("tile_*.npz")):
        xyz = np.load(q)["xyz"]
        r = np.hypot(xyz[:, None, 0] - pos[None, :, 0], xyz[:, None, 1] - pos[None, :, 1]) - rad[None]
        trunk = (np.abs(r) < 0.01).any(1) & (xyz[:, 2] > 0.03 * xyz[:, 0] + 0.3)
        sem = np.where(trunk, 1, 2).astype(np.int8)
        sem[::97] = -1
        np.savez(d / "sem0.02" / q.name, sem=sem, hag=(xyz[:, 2] - 0.03 * xyz[:, 0]).astype(np.float16))
    fj = d / "frame.json"
    fr = json.loads(fj.read_text()); fr["semantic"] = dict(stamp={"synthetic": 1}, tally={}); fj.write_text(json.dumps(fr))


def test_a_longer_old_run_is_moved_aside_not_continued():
    """Run 1 (150,000-step schedule, NaN at ~18,850) must not be resumed by a
    50,000-step run: its files move aside and training starts from 0."""
    with tempfile.TemporaryDirectory() as t:
        env, pos, rad = _build(Path(t))
        env.update(V2_ENCODER="fake", V2_DEVICE="cpu", V2_STEPS="12", V2_EVAL_EVERY="12", V2_ACCUM="1",
                   V2_CAP="40000", V2_PROGRESS_S="0")
        r = subprocess.run([sys.executable, str(V2 / "train.py")], env=env, capture_output=True, text=True)
        assert r.returncode == 0, r.stdout[-2000:] + r.stderr[-2000:]
        env["V2_STEPS"] = "6"; env["V2_EVAL_EVERY"] = "6"
        r = subprocess.run([sys.executable, str(V2 / "train.py")], env=env, capture_output=True, text=True)
        assert r.returncode == 0, r.stdout[-2000:] + r.stderr[-2000:]
        assert "12-step schedule, this one 6: not continuing it" in r.stdout and "resuming" not in r.stdout, r.stdout[-2000:]
        models, res = Path(t) / "data" / "stemnet_v2" / "models", Path(t) / "res"
        for q in (models / "head_last_12steps.pth", models / "head_best_12steps.pth", res / "stemnet_v2_train_12steps.txt"):
            assert q.exists(), q
        log = (res / "stemnet_v2_train.txt").read_text()
        assert "done: 6 steps" in log and "step 1/6" in log, log[-1500:]


def test_tree_task_trains_evaluates_and_keeps_its_own_checkpoints():
    import torch
    with tempfile.TemporaryDirectory() as t:
        env, pos, rad = _build(Path(t))
        _add_sem(t, pos, rad)
        r = subprocess.run([sys.executable, str(V2 / "prep_cache.py")], env=env, capture_output=True, text=True)
        assert r.returncode == 0 and "tree / non-tree labels: included" in r.stdout, r.stdout[-1500:] + r.stderr[-1500:]
        _with_env(env)
        import data as D
        a, b = D.TileStore(), D.CacheStore()
        ca, cb = a.crop(24.0, 15.0, 8.0), b.crop(24.0, 15.0, 8.0)
        oa, ob = np.lexsort(ca["xyz"].T), np.lexsort(cb["xyz"].T)
        assert np.array_equal(ca["sem"][oa], cb["sem"][ob]) and (cb["sem"] == 1).any()
        env.update(V2_ENCODER="fake", V2_DEVICE="cpu", V2_STEPS="12", V2_EVAL_EVERY="6", V2_ACCUM="1",
                   V2_CAP="40000", V2_PROGRESS_S="0", V2_TASK="base+tree")
        r = subprocess.run([sys.executable, str(V2 / "train.py")], env=env, capture_output=True, text=True)
        assert r.returncode == 0, r.stdout[-3000:] + r.stderr[-3000:]
        res, models = Path(t) / "res", Path(t) / "data" / "stemnet_v2" / "models"
        log = (res / "stemnet_v2_train_tree.txt").read_text()
        assert log.count("TREE all labelled: mIoU") == 2 and log.count("TREE matched only") == 2, log[-2000:]
        assert " tree " in log and "done: 12 steps" in log and "best tree mIoU" in log, log[-1500:]
        assert not (res / "stemnet_v2_train.txt").exists()            # base-task log untouched
        ck = torch.load(models / "head_tree_last.pth", weights_only=False)
        assert "tree.weight" in ck["head"] and ck["task"] == "base+tree"
        assert (models / "head_tree_best_miou.pth").exists() and not (models / "head_last.pth").exists()
        assert (res / "51_stemnet_v2_val_tree.png").stat().st_size > 10000
        env["V2_STEPS"] = "14"
        r = subprocess.run([sys.executable, str(V2 / "train.py")], env=env, capture_output=True, text=True)
        assert r.returncode == 0 and "resuming from step 12" in r.stdout, r.stdout[-2000:] + r.stderr[-2000:]


def test_tree_task_refuses_crops_without_tree_labels():
    with tempfile.TemporaryDirectory() as t:
        env, pos, rad = _build(Path(t))
        env.update(V2_ENCODER="fake", V2_DEVICE="cpu", V2_STEPS="2", V2_TASK="base+tree")
        r = subprocess.run([sys.executable, str(V2 / "train.py")], env=env, capture_output=True, text=True)
        assert r.returncode != 0 and "run v2_semantic" in r.stdout + r.stderr, r.stdout[-1500:] + r.stderr[-1500:]


def test_blocks_cover_every_point_once_and_find_every_base_once():
    """Perfect predictions through v2_infer's blocks: no point written twice or
    lost at a block edge, no base lost or doubled where blocks meet."""
    with tempfile.TemporaryDirectory() as t:
        env, pos, rad = _build(Path(t))
        _add_sem(t, pos, rad)
        _with_env(env)
        import data as D, infer as I, vote as V
        store = D.TileStore()
        occ = np.load(Path(t) / "data" / "stemnet_v2" / "regions.npz")["occupancy"]
        cores = I.blocks((0.0, W, 0.0, W), occ, block=7.0)              # 5 x 5 blocks, edges cut through trees
        oracle = lambda c: (c["gate"].astype(np.float32), c["offset"].astype(np.float32),  # noqa: E731
                            (c["sem"] == 1).astype(np.float32))
        params = dict(p_thresh=V.P_THRESH, peak_sep=V.PEAK_SEP, peak_min=V.PEAK_MIN)
        keys, det = [], []
        for core, R, inside, d, s, n in I.iter_blocks(oracle, store, cores, params, np.random.default_rng(0), margin=4.0):
            keys.append(R["tile"][inside] * 10**8 + R["row"][inside]); det.append(d)
            assert np.array_equal(R["tree"][inside] >= 0.5, R["sem"][inside] == 1)
        keys, det = np.concatenate(keys), np.concatenate(det)
        total = sum(len(store.tile(ij)["xyz"]) for ij in store.have)
        assert len(keys) == len(np.unique(keys)) == total, (len(keys), len(np.unique(keys)), total)
        m = V.match(det, np.ones(len(det)), pos, 0.3)
        assert m["recall"] == 1.0 and m["precision"] == 1.0, {k: v for k, v in m.items() if k != "pair"}


def test_validate_and_infer_run_end_to_end():
    import laspy
    with tempfile.TemporaryDirectory() as t:
        env, pos, rad = _build(Path(t))
        _add_sem(t, pos, rad)
        d = Path(t) / "data" / "stemnet_v2"
        split = np.zeros((int(W), int(W)), np.uint8)
        split[:, 10:12] = 3; split[:, 12:18] = 2; split[:, 18:20] = 3; split[:, 20:] = 1   # train | test | val
        np.savez(d / "regions.npz", split=split, occupancy=np.ones_like(split, np.int32))
        r = subprocess.run([sys.executable, str(V2 / "prep_cache.py")], env=env, capture_output=True, text=True)
        assert r.returncode == 0, r.stdout[-1500:] + r.stderr[-1500:]
        env.update(V2_ENCODER="fake", V2_DEVICE="cpu", V2_STEPS="12", V2_EVAL_EVERY="12", V2_ACCUM="1",
                   V2_CAP="40000", V2_PROGRESS_S="0", V2_TASK="base+tree", V2_INFER_BLOCK="12",
                   V2_VAL_P="0.5,0.9", V2_VAL_SEP="0.3", V2_VAL_MASS="1,10")
        for script in ("train.py", "validate.py", "infer.py"):
            r = subprocess.run([sys.executable, str(V2 / script)], env=env, capture_output=True, text=True)
            assert r.returncode == 0, (script, r.stdout[-3000:] + r.stderr[-3000:])
        res = Path(t) / "res"
        log = (res / "stemnet_v2_validate.txt").read_text()
        for s in ("checkpoint head_tree_best.pth", "VAL default settings", "sweep on val: 4 settings",
                  "TEST default settings", "TEST val-tuned settings", "TREE test all labelled", "recall@0.5 by tier"):
            assert s in log, (s, log[-3000:])
        prm = json.loads((d / "models" / "detect_params.json").read_text())
        assert prm["checkpoint"] == "head_tree_best.pth" and prm["p_thresh"] in (0.5, 0.9)
        for f in ("stemnet_v2_test_detections.csv", "53_stemnet_v2_test.png", "54_stemnet_v2_test_tree.png"):
            assert (res / f).exists(), f
        log = (res / "stemnet_v2_infer.txt").read_text()
        assert "detection settings tuned on val" in log and "9 blocks" in log, log[:2000]
        assert "tree head vs the campus_pred_B24 labels" in log and "  test " in log, log[-2500:]
        assert (res / "52_stemnet_v2_infer_map.png").exists() and (res / "stemnet_v2_infer_bases.csv").exists()
        lazs = sorted((d / "infer" / "head_tree_best").glob("block_*.laz"))
        assert len(lazs) == 9, lazs
        got = [laspy.read(str(q)) for q in lazs]
        xyz = np.concatenate([np.column_stack([g.x, g.y, g.z]) for g in got])
        grid = np.concatenate([np.load(q)["xyz"] for q in sorted((d / "grid0.02").glob("tile_*.npz"))]).astype(np.float64)
        assert len(xyz) == len(grid), (len(xyz), len(grid))
        from scipy.spatial import cKDTree
        dist, idx = cKDTree(grid).query(xyz)                         # each grid point once, at its place (mm scale)
        assert dist.max() < 0.001 and len(np.unique(idx)) == len(grid), (dist.max(), len(np.unique(idx)))
        tp = np.concatenate([np.asarray(g.tree_prob) for g in got])
        assert ((tp >= 0) & (tp <= 1)).all()
        assert set(np.unique(np.concatenate([np.asarray(g.split) for g in got]))) == {0, 1, 2, 3}
        assert np.array_equal(np.concatenate([np.asarray(g.pred_tree) for g in got]), (tp >= 0.5).astype(np.uint8))
        r = subprocess.run([sys.executable, str(V2 / "validate.py")], env=env, capture_output=True, text=True)
        assert r.returncode == 0 and r.stdout.count("saved predictions reused") == 2, r.stdout[-2000:] + r.stderr[-2000:]
        r = subprocess.run([sys.executable, str(V2 / "curves.py")], env=dict(env, V2_CURVE_P="0.5,0.9"),
                           capture_output=True, text=True)
        assert r.returncode == 0, r.stdout[-3000:] + r.stderr[-3000:]
        log = (res / "stemnet_v2_curves.txt").read_text()
        assert "TEST (held out): AP@0.5 m" in log and "chosen on val" in log and "v2_infer's settings on test" in log, log
        assert (res / "57_stemnet_v2_curves.png").exists() and (res / "stemnet_v2_curves.csv").exists()
        r = subprocess.run([sys.executable, str(V2 / "inspect_infer.py")], env=dict(env, V2_INSPECT_SIZE="6", V2_INSPECT_SPOTS="2"),
                           capture_output=True, text=True)
        assert r.returncode == 0, r.stdout[-3000:] + r.stderr[-3000:]
        log = (res / "stemnet_v2_inspect.txt").read_text()
        assert f"{len(grid):,} of {len(grid):,} points (100.00%)" in log, log[-3000:]       # every LAZ point found in the grid
        assert "surveyed trunks, train" in log and "test detections" in log and "calibration" in log, log[-3000:]
        assert (res / "55_stemnet_v2_inspect.png").exists() and list(res.glob("56_stemnet_v2_inspect_tree_f*_1.png")), log[-2000:]


def test_compare_scores_both_models_into_one_spreadsheet():
    from openpyxl import load_workbook
    with tempfile.TemporaryDirectory() as t:
        env, pos, rad = _build(Path(t))
        _add_sem(t, pos, rad)
        d = Path(t) / "data" / "stemnet_v2"
        split = np.zeros((int(W), int(W)), np.uint8)
        split[:, 10:12] = 3; split[:, 12:18] = 2; split[:, 18:20] = 3; split[:, 20:] = 1   # train | test | val
        np.savez(d / "regions.npz", split=split, occupancy=np.ones_like(split, np.int32))
        r = subprocess.run([sys.executable, str(V2 / "prep_cache.py")], env=env, capture_output=True, text=True)
        assert r.returncode == 0, r.stdout[-1500:] + r.stderr[-1500:]
        env.update(V2_ENCODER="fake", V2_DEVICE="cpu", V2_STEPS="12", V2_EVAL_EVERY="12", V2_ACCUM="1",
                   V2_CAP="40000", V2_PROGRESS_S="0", V2_INFER_BLOCK="12", V2_CURVE_P="0.5,0.9")
        for task in ("base", "base+tree"):
            r = subprocess.run([sys.executable, str(V2 / "train.py")], env=dict(env, V2_TASK=task), capture_output=True, text=True)
            assert r.returncode == 0, (task, r.stdout[-2000:] + r.stderr[-2000:])
        r = subprocess.run([sys.executable, str(V2 / "compare.py")], env=env, capture_output=True, text=True)
        assert r.returncode == 0, r.stdout[-3000:] + r.stderr[-3000:]
        res = Path(t) / "res"
        log = (res / "stemnet_v2_compare.txt").read_text()
        assert log.count("val-chosen cut (vote mass") == 2 and "base-only" in log and "base+tree" in log, log[-3000:]
        assert "recall at precision >=" in log and "at maximum recall" in log, log[-3000:]
        wb = load_workbook(res / "stemnet_v2_compare.xlsx", read_only=True)
        assert wb.sheetnames == ["README", "Summary", "PerGate", "Breakdown", "Curves", "Detections", "SurveyTrees"]
        rows = {n: list(wb[n].iter_rows(values_only=True)) for n in wb.sheetnames}
        assert [r_[0] for r_ in rows["Summary"][1:]] == ["base-only", "base+tree"], rows["Summary"]
        assert len(rows["PerGate"]) == 1 + 2 * 2 * 2                         # models x splits x gates
        n_gt = sum(1 for x, y in pos if 12 <= x < 18) + sum(1 for x, y in pos if x >= 20)
        assert len(rows["SurveyTrees"]) == 1 + n_gt and len(rows["SurveyTrees"][0]) == 10 + 2 * 4
        assert (res / "58_stemnet_v2_compare.png").exists()
        r = subprocess.run([sys.executable, str(V2 / "compare.py")], env=env, capture_output=True, text=True)
        assert r.returncode == 0 and r.stdout.count("saved predictions reused") == 4, r.stdout[-2000:] + r.stderr[-2000:]
        assert (res / "stemnet_v2_compare.txt").read_text().splitlines()[-3:] == log.splitlines()[-3:]
        import laspy
        r = subprocess.run([sys.executable, str(V2 / "export_val.py")], env=env, capture_output=True, text=True)
        assert r.returncode == 0, r.stdout[-3000:] + r.stderr[-3000:]
        las = laspy.read(str(res / "stemnet_v2_val.laz"))
        mk = np.asarray(las.marker)
        assert np.array_equal(mk, np.asarray(las.classification))
        grid = np.concatenate([np.load(q)["xyz"] for q in sorted((d / "grid0.02").glob("tile_*.npz"))])
        val = grid[grid[:, 0] >= 20]
        assert (mk == 0).sum() == len(val), ((mk == 0).sum(), len(val))      # every val point once, in the same file
        assert set(np.unique(mk)) == {0, 1, 2, 3}                           # + base-only, base+tree, survey posts
        assert "scan RGB joined for" in r.stdout and "(100.00%)" in r.stdout, r.stdout[-2000:]
        for f in ("gate_base_only", "gate_base_tree", "vote_base_only", "vote_base_tree", "tree_prob", "pred_tree",
                  "label_tree", "tree_vs_b24", "base_id", "vote_mass", "matched"):
            assert f in las.point_format.extra_dimension_names, f
        bid = np.asarray(las.base_id)
        n_det = int(bid[mk == 2].max()) + 1
        vid = np.asarray(las.vote_base_tree)
        assert set(np.unique(vid[vid >= 0])) == set(range(n_det)) and (vid[mk > 0] == -1).all()   # every detection has its voters
        assert (bid[mk == 0] == -1).all() and (np.asarray(las.red)[mk == 2] == 255 * 257).all()


def test_detection_curve_equals_rematching_at_every_threshold():
    """One greedy matching, read off cumulatively, must equal matching only the
    detections above each threshold; AP is 1 for a perfect ranking."""
    sys.path.insert(0, str(V2))
    import curves as C, vote as V
    rng = np.random.default_rng(3)
    gt = rng.uniform(0, 30, (60, 2))
    det = np.r_[gt[:45] + rng.normal(0, 0.25, (45, 2)), rng.uniform(0, 30, (40, 2)), gt[:10] + rng.normal(0, 0.3, (10, 2))]
    score = rng.uniform(1, 100, len(det))
    ignore = det[:, 0] > 27
    for radius in (0.5, 0.3):
        c = C.curve(det, score, gt, radius, ignore)
        for t in np.r_[np.sort(score)[::7], 0.0]:
            k = score >= t
            m = V.match(det[k], score[k], gt, radius, ignore=ignore[k])
            n = np.flatnonzero(c["score"] >= t)
            tp, fp = (c["tp"][n[-1]], c["fp"][n[-1]]) if len(n) else (0, 0)
            assert (tp, fp) == (m["tp"], m["fp"]), (radius, t, tp, fp, m["tp"], m["fp"])
    perfect = C.curve(gt.copy(), np.arange(len(gt), 0, -1.0), gt, 0.5)
    assert abs(C.ap(perfect) - 1.0) < 1e-9
    worse = C.curve(np.r_[gt + 5.0, gt], np.r_[np.full(len(gt), 9.0), np.full(len(gt), 1.0)], gt, 0.5)
    assert C.ap(worse) < 0.6 and abs(worse["tp"][-1] / len(gt) - 1.0) < 1e-9


def test_kmeans_recovers_separated_blobs():
    import torch
    sys.path.insert(0, str(V2))
    import segment_probe as S
    rng = np.random.default_rng(0)
    truth = np.repeat(np.arange(4), 500)
    X = torch.from_numpy((rng.normal(0, 0.3, (2000, 6)) + np.eye(4, 6)[truth] * 5).astype(np.float32))
    C, inertia, it = S.kmeans(X, 4, rng)
    lab = ((X * X).sum(1, keepdim=True) - 2 * X @ C.T + (C * C).sum(1)[None]).argmin(1).numpy()
    for t in range(4):                                # each blob is exactly one cluster
        assert len(set(lab[truth == t])) == 1, t
    assert len(set(lab)) == 4


def test_segment_probe_encodes_every_point_in_the_region():
    with tempfile.TemporaryDirectory() as t:
        env, pos, rad = _build(Path(t))
        env.update(V2_ENCODER="fake", V2_DEVICE="cpu", V2_SEG_SIZE="8", V2_SEG_XY="25,15", V2_SEG_K="2,4",
                   V2_SEG_FIT="20000", V2_CAP="20000")         # cap forces crops to be split
        r = subprocess.run([sys.executable, str(V2 / "segment_probe.py")], env=env, capture_output=True, text=True)
        assert r.returncode == 0, r.stdout[-3000:] + r.stderr[-3000:]
        res = Path(t) / "res"
        log = (res / "stemnet_v2_segment.txt").read_text()
        assert "over the 20,000 cap" in log and "summary" in log, log[-1500:]
        for v in ("sonata", "sonata+feat", "sonata+feat_bal"):
            assert f"== {v} " in log
            ply = (res / f"stemnet_v2_segment_{v}.ply").read_bytes()
            n = int(ply.split(b"element vertex ")[1].split(b"\n")[0])
            _with_env(env)
            import data as D
            store = D.TileStore()
            allxyz = np.concatenate([store.tile(ij)["xyz"] for ij in sorted(store.have)])
            want = ((allxyz[:, 0] >= 21) & (allxyz[:, 0] < 29) & (allxyz[:, 1] >= 11) & (allxyz[:, 1] < 19)).sum()
            assert n == want, (v, n, want)
        for p in ("47_stemnet_v2_segment_top.png", "48_stemnet_v2_segment_side.png"):
            assert (res / p).stat().st_size > 10000, p


def test_view_squares_draws_every_point():
    with tempfile.TemporaryDirectory() as t:
        env, pos, rad = _build(Path(t))
        env.update(V2_VIEW_SIZE="8", V2_VIEW_N="2", V2_VIEW_XY="25,15")
        r = subprocess.run([sys.executable, str(V2 / "view_squares.py")], env=env, capture_output=True, text=True)
        assert r.returncode == 0, r.stdout[-3000:] + r.stderr[-3000:]
        res = Path(t) / "res"
        log = (res / "stemnet_v2_view.txt").read_text()
        assert "3 squares" in log, log
        _with_env(env)
        import view_squares as VS
        xyz, rgb = VS.load_square(25.0, 15.0)
        _with_env(env)
        import data as D
        allxyz = np.concatenate([D.TileStore().tile(ij)["xyz"] for ij in sorted(D.TileStore().have)])
        want = ((np.abs(allxyz[:, 0] - 25) < 4) & (np.abs(allxyz[:, 1] - 15) < 4)).sum()
        assert len(xyz) == want and f"{want:,} points" in log, (len(xyz), want, log)
        top, bot, col, empty = VS.raster(xyz, rgb, 21.0, 11.0)      # a pixel is empty iff no point is in it
        i = np.clip(((xyz[:, 0] - 21) / VS.PX).astype(int), 0, top.shape[0] - 1)
        j = np.clip(((xyz[:, 1] - 11) / VS.PX).astype(int), 0, top.shape[0] - 1)
        hit = np.zeros_like(empty); hit[j, i] = True
        assert np.array_equal(hit, ~empty)
        for k in range(3):
            assert (res / f"49_stemnet_v2_view_{k}.png").stat().st_size > 10000, k
        assert (res / "49_stemnet_v2_view_map.png").exists()


if __name__ == "__main__":
    fails = 0
    for name, fn in list(globals().items()):
        if name.startswith("test_") and callable(fn):
            try:
                fn(); print(f"PASS {name}")
            except Exception as e:                    # noqa: BLE001
                fails += 1; print(f"FAIL {name}: {type(e).__name__}: {e}")
    print(f"\n{'All' if not fails else fails} StemNet v2 train test(s) {'passed' if not fails else 'FAILED'}.")
    sys.exit(1 if fails else 0)
